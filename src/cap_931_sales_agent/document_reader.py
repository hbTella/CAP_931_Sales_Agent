
from io import BytesIO

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


def read_uploaded_document(uploaded_file) -> str:
    """Extract text from an uploaded PDF, TXT, DOCX, or PPTX file."""

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="replace").strip()

    if file_name.endswith(".pdf"):
        pdf = PdfReader(BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or "" for page in pdf.pages
        ).strip()

    if file_name.endswith(".docx"):
        document = Document(BytesIO(file_bytes))
        return "\n".join(
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ).strip()

    if file_name.endswith(".pptx"):
        presentation = Presentation(BytesIO(file_bytes))
        slide_text = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    slide_text.append(shape.text.strip())

        return "\n".join(slide_text).strip()

    raise ValueError("Unsupported file type.")
