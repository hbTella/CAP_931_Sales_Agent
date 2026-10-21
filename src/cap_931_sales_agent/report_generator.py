import re
from datetime import date

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def insert_after(anchor, text, style="Normal"):
    """Insert a formatted paragraph after an existing paragraph."""
    new_element = OxmlElement("w:p")
    anchor._p.addnext(new_element)

    paragraph = Paragraph(new_element, anchor._parent)

    # Use the requested style when the template supports it.
    try:
        paragraph.style = style
    except KeyError:
        paragraph.style = "Normal"

    paragraph.add_run(text)

    return paragraph


def add_research_text(anchor, research):
    """Convert research text into readable Word paragraphs."""

    # Split the text into lines and separate questions that share a line.
    research = re.sub(
        r"\s+(?=(?:Q[1-5]\.?|[a-e]\.)\s*)",
        "\n",
        research,
        flags=re.IGNORECASE,
    )

    for line in research.splitlines():
        line = line.strip()

        if not line:
            continue

        # Convert Markdown links to readable text and URLs.
        line = re.sub(
            r"\[([^\]]+)\]\((https?://[^)]+)\)",
            r"\1 — \2",
            line,
        )

        # Convert Markdown headings into Word headings.
        if line.startswith("#"):
            heading = re.sub(r"^#+\s*", "", line)
            anchor = insert_after(anchor, heading, "Heading 3")

        # Convert bullet points into Word paragraphs.
        elif line.startswith(("- ", "* ")):
            anchor = insert_after(anchor, f"• {line[2:]}", "Normal")

        # Convert numbered section headings into Word headings.
        elif re.match(r"^\d+\.\s+[A-Z]", line):
            heading = re.sub(r"^\d+\.\s+", "", line)
            anchor = insert_after(anchor, heading, "Heading 3")

        # Give each question its own paragraph.
        elif re.match(r"^(?:Q[1-5]\.?|[a-e]\.)\s*", line, re.IGNORECASE):
            anchor = insert_after(anchor, line, "Normal")

        # Keep other content in its own paragraph.
        else:
            anchor = insert_after(anchor, line, "Normal")

    return anchor


def generate_sales_report(
    template_path: str,
    output_path: str,
    product_name: str,
    product_category: str,
    value_proposition: str,
    target_customer: str,
    company_url: str,
    company_research: str,
    competitor_research: list[dict],
    product_overview: str = "",
) -> str:
    """Fill the Per Scholas Word template with sales research."""

    document = Document(template_path)

    # Find the template's main section headings and placeholders.
    paragraphs = document.paragraphs
    overview_placeholder = None
    steps_placeholder = None
    conclusion_placeholder = None
    appendix_heading = None

    for paragraph in paragraphs:
        text = paragraph.text.strip()

        if text == "[Title]":
            paragraph.text = f"CAP 931 Sales Research: {product_name}"

        elif "[Learner Name]" in paragraph.text:
            paragraph.text = (
                f"Learner Name: Hima Bindu Tella   \n"    
                f"Date: {date.today().strftime('%B %d, %Y')}"
            )

        elif text.startswith("[Create a short summary"):
            overview_placeholder = paragraph

        elif text.startswith("[Detail the steps"):
            steps_placeholder = paragraph

        elif text.startswith("[Add any takeaways"):
            conclusion_placeholder = paragraph

        elif text == "Appendix":
            appendix_heading = paragraph

    # Fill in the Overview section.
    if overview_placeholder is not None:
        overview_placeholder.text = (
            f"This report examines {company_url} as a potential sales "
            f"prospect for {product_name}. The product category is "
            f"{product_category or 'not provided'}. The value proposition "
            f"is {value_proposition.rstrip('. ') if value_proposition else 'not provided'}. "
            f"The target customer "
            f"is {target_customer.rstrip('. ') if target_customer else 'not provided'}."
        )

    # Fill the Steps section with properly formatted research.
    if steps_placeholder is not None:
        steps_placeholder.text = (
            "The following steps were completed to prepare for sales discovery:"
        )
        
        steps_placeholder.style = "Normal"

        # Remove numbering and indentation from the introductory paragraph.
        pPr = steps_placeholder._p.get_or_add_pPr()
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)

        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            pPr.remove(ind)

        anchor = insert_after(
            steps_placeholder,
            "Step 1: Gather product and target-customer information.",
        )
        # Add the uploaded product overview to the report.
        if product_overview.strip():
            anchor = insert_after(
                anchor,
                "Uploaded Product Overview",
                "Heading 2",
            )
            anchor = add_research_text(anchor, product_overview)

        anchor = insert_after(
            anchor,
            "Step 2: Research the prospective company's public website.",
        )
        anchor = insert_after(
            anchor,
            "Company Research Findings",
            "Heading 2",
        )
        anchor = add_research_text(anchor, company_research)

        anchor = insert_after(
            anchor,
            "Step 3: Research the supplied competitor websites.",
        )
        anchor = insert_after(
            anchor,
            "Competitor Research Findings",
            "Heading 2",
        )

        if competitor_research:
            for item in competitor_research:
                anchor = insert_after(
                    anchor,
                    f"Competitor: {item['url']}",
                    "Heading 3",
                )
                anchor = add_research_text(anchor, item["summary"])
        else:
            anchor = insert_after(
                anchor,
                "No competitor research was provided.",
            )

    
    # Fill in the Conclusion section.
    if conclusion_placeholder is not None:
        if competitor_research:
            conclusion_text = (
                "The research identifies company priorities, competitor "
                "positioning, and questions for a sales discovery meeting. "
                "The sales representative should validate these findings "
                "with the prospect before making recommendations."
            )
        else:
            conclusion_text = (
                "The research identifies company priorities and questions "
                "for a sales discovery meeting. Competitor research was not "
                "provided, so competitor positioning was not evaluated. "
                "The sales representative should validate these findings "
                "with the prospect before making recommendations."
            )

        conclusion_placeholder.text = conclusion_text


    # Add a clean source list under Appendix.
    if appendix_heading is not None:
        source_text = [company_url]

        for item in competitor_research:
            source_text.append(item["url"])
            source_text.extend(
                re.findall(r"https?://[^\s)\]]+", item["summary"])
            )

        source_text.extend(
            re.findall(r"https?://[^\s)\]]+", company_research)
        )

        # Remove duplicate URLs and trailing punctuation.
        sources = list(
            dict.fromkeys(url.rstrip(".,;") for url in source_text)
        )

        anchor = appendix_heading

        for url in sources:
            anchor = insert_after(anchor, url, "List Bullet")

    # Save the finished report without modifying the template.
    document.save(output_path)

    return output_path
