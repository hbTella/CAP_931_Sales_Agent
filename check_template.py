from docx import Document

# Open the Per Scholas template
document = Document("per_scholas_template.docx")

# Display the document's paragraphs
print("=== DOCUMENT PARAGRAPHS ===")
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text.strip():
        print(f"Paragraph {index}: {paragraph.text}")

# Check whether the template contains tables
print("\n=== DOCUMENT TABLES ===")
for table_index, table in enumerate(document.tables):
    print(f"Table {table_index}:")
    for row in table.rows:
        print(" | ".join(cell.text for cell in row.cells))

# Inspect headers and footers
print("\n=== HEADERS AND FOOTERS ===")
for section_index, section in enumerate(document.sections):
    print(f"Section {section_index} header:")
    for paragraph in section.header.paragraphs:
        if paragraph.text.strip():
            print(paragraph.text)

    print(f"Section {section_index} footer:")
    for paragraph in section.footer.paragraphs:
        if paragraph.text.strip():
            print(paragraph.text)